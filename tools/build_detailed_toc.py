from __future__ import annotations

import argparse
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT_NAME = "TH SarabunPSK"
THAI_PAGE_LETTERS = list("กขคงจฉชซฌญฎฏฐฑฒณดตถทธนบปผฝพฟภมยรลวศษสหฬอฮ")


@dataclass
class Entry:
    title: str
    page: str
    level: int = 0


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\x0b", " ")).strip()


def set_run_font(run, size: float = 16, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)


def set_paragraph_base(paragraph, *, alignment=None, left_inches: float = 0.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.keep_together = True
    fmt.widow_control = False
    if alignment is not None:
        paragraph.alignment = alignment
    if left_inches:
        fmt.left_indent = Inches(left_inches)
    tabs = fmt.tab_stops
    tabs.add_tab_stop(Inches(5.73), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)


def add_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_base(p, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run(title), size=18, bold=True)


def add_column_header(doc: Document, left_label: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_base(p)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(f"{left_label}\tหน้า"), size=16)


def add_entry(doc: Document, entry: Entry) -> None:
    p = doc.add_paragraph()
    indent = {0: 0.0, 1: 0.27, 2: 0.58}.get(entry.level, 0.58)
    set_paragraph_base(p, left_inches=indent)
    set_run_font(p.add_run(f"{entry.title}\t{entry.page}"), size=16)


def estimate_units(entry: Entry) -> int:
    capacity = {0: 61, 1: 57, 2: 52}.get(entry.level, 52)
    return max(1, math.ceil((len(entry.title) + 4) / capacity))


def paginate(entries: list[Entry], capacity: int = 29) -> list[list[Entry]]:
    pages: list[list[Entry]] = []
    current: list[Entry] = []
    used = 0
    for entry in entries:
        units = estimate_units(entry)
        if entry.level == 0 and current and used + units + 1 > capacity:
            pages.append(current)
            current = []
            used = 0
        if current and used + units > capacity:
            pages.append(current)
            current = []
            used = 0
        current.append(entry)
        used += units
    if current:
        pages.append(current)
    return pages


def thai_label_for_physical_page(physical_page: int) -> str:
    index = 6 + physical_page - 1
    if index >= len(THAI_PAGE_LETTERS):
        raise ValueError(f"Thai preliminary page exceeds supported sequence: {physical_page}")
    return THAI_PAGE_LETTERS[index]


def chapter_entries(outline: dict) -> list[Entry]:
    entries: list[Entry] = []
    chapter_number = 0
    for item in outline["items"]:
        style = str(item.get("style", ""))
        text = clean_text(str(item.get("text", "")))
        number = clean_text(str(item.get("list_string", "")))
        page = str(int(item["absolute_page"]))
        if not text:
            continue
        if style.startswith("Heading 1"):
            chapter_number += 1
            entries.append(Entry(f"บทที่ {chapter_number}  {text}", page, 0))
        elif style.startswith("Heading 2"):
            entries.append(Entry(f"{number} {text}".strip(), page, 1))
        elif style.startswith("Heading 3"):
            include = bool(re.match(r"^(?:[2-5]\.|1\.5\.)", number)) and len(text) <= 100
            if include:
                entries.append(Entry(f"{number} {text}".strip(), page, 2))
    return entries


def appendix_entries(outline: dict, offset: int) -> list[Entry]:
    entries: list[Entry] = []
    title_added = False
    for item in outline["items"]:
        text = clean_text(str(item.get("text", "")))
        style = str(item.get("style", ""))
        page = str(offset + int(item["absolute_page"]))
        if text.startswith("ภาคผนวก ก") and not title_added:
            entries.append(Entry(text, page, 0))
            title_added = True
        elif "Appendix Heading" in style:
            entries.append(Entry(text, page, 1))
    return entries


def caption_entries(outline: dict, prefix: str) -> list[Entry]:
    rows = []
    for item in outline["items"]:
        text = clean_text(str(item.get("text", "")))
        if text.startswith(prefix):
            rows.append((int(item["absolute_page"]), int(item.get("index", 0)), text))

    if prefix == "รูปที่":
        rows.extend(
            [
                (15, 5658, "รูปที่ 2.9 โมดูลสวิตช์โหลด HW-548 ที่ใช้ควบคุมการจ่ายไฟให้เซนเซอร์และโมดูล GPS (ที่มา: ผู้จัดทำ)"),
                (17, 6870, "รูปที่ 2.11 โมดูลควบคุมการชาร์จ CN3791 (ที่มา: ผู้จัดทำ)"),
            ]
        )

    rows.sort(key=lambda value: (value[0], value[1]))
    return [Entry(text, str(page), 0) for page, _index, text in rows]


def clear_container(container) -> None:
    element = container._element
    for child in list(element):
        element.remove(child)
    element.append(OxmlElement("w:p"))


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=14)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE \\* MERGEFORMAT "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "ช"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, cached, end):
        run._element.append(node)


def configure_document(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.55)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = False

    clear_container(section.header)
    header_p = section.header.paragraphs[0]
    add_page_field(header_p)
    clear_container(section.footer)

    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), "thaiLetters")
    pg_num.set(qn("w:start"), "7")

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def render_grouped_section(doc: Document, title: str, left_label: str, pages: list[list[Entry]]) -> None:
    for page_index, entries in enumerate(pages):
        if len(doc.paragraphs) > 0:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        shown_title = title if page_index == 0 else f"{title} (ต่อ)"
        add_title(doc, shown_title)
        add_column_header(doc, left_label)
        for entry in entries:
            add_entry(doc, entry)


def build(args) -> None:
    template = Path(args.template)
    output = Path(args.output)
    combined = json.loads(Path(args.combined_outline).read_text(encoding="utf-8"))
    appendix = json.loads(Path(args.appendix_outline).read_text(encoding="utf-8"))

    main_entries = chapter_entries(combined)
    bibliography_start = int(combined["page_count"]) + 1
    bibliography_pages = int(args.bibliography_pages)
    appendix_start = bibliography_start + bibliography_pages
    appendix_page_count = int(appendix["page_count"])
    biography_start = appendix_start + appendix_page_count

    back_entries = [Entry("บรรณานุกรม", str(bibliography_start), 0)]
    back_entries.extend(appendix_entries(appendix, appendix_start - 1))
    back_entries.append(Entry("ประวัติผู้เขียน", str(biography_start), 0))

    table_entries = caption_entries(combined, "ตารางที่")
    figure_entries = caption_entries(combined, "รูปที่")

    placeholder_front = [
        Entry("บทคัดย่อ", "ง", 0),
        Entry("กิตติกรรมประกาศ", "ฉ", 0),
        Entry("สารบัญตาราง", "X", 0),
        Entry("สารบัญรูป", "X", 0),
        Entry("คำอธิบายสัญลักษณ์และคำย่อ", "X", 0),
    ]

    toc_pages = paginate(placeholder_front + main_entries + back_entries, capacity=22)
    table_pages = paginate(table_entries, capacity=29)
    figure_pages = paginate(figure_entries, capacity=31)

    table_physical = len(toc_pages) + 1
    figure_physical = len(toc_pages) + len(table_pages) + 1
    symbols_physical = len(toc_pages) + len(table_pages) + len(figure_pages) + 1

    front_entries = [
        Entry("บทคัดย่อ", "ง", 0),
        Entry("กิตติกรรมประกาศ", "ฉ", 0),
        Entry("สารบัญตาราง", thai_label_for_physical_page(table_physical), 0),
        Entry("สารบัญรูป", thai_label_for_physical_page(figure_physical), 0),
        Entry("คำอธิบายสัญลักษณ์และคำย่อ", thai_label_for_physical_page(symbols_physical), 0),
    ]
    toc_pages = paginate(front_entries + main_entries + back_entries, capacity=22)

    doc = Document(str(template))
    configure_document(doc)

    # render_grouped_section inserts a page break when the document already has
    # a paragraph; remove the empty body paragraph first by rendering page one
    # directly, then use explicit page breaks for later pages/sections.
    for page_index, entries in enumerate(toc_pages):
        if page_index:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        add_title(doc, "สารบัญ" if page_index == 0 else "สารบัญ (ต่อ)")
        add_column_header(doc, "หัวข้อ")
        for entry in entries:
            add_entry(doc, entry)

    for title, left_label, pages in (
        ("สารบัญตาราง", "ตาราง", table_pages),
        ("สารบัญรูป", "รูป", figure_pages),
    ):
        for page_index, entries in enumerate(pages):
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            add_title(doc, title if page_index == 0 else f"{title} (ต่อ)")
            add_column_header(doc, left_label)
            for entry in entries:
                add_entry(doc, entry)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))

    print(json.dumps({
        "output": str(output),
        "toc_pages_planned": len(toc_pages),
        "table_pages_planned": len(table_pages),
        "figure_pages_planned": len(figure_pages),
        "total_pages_planned": len(toc_pages) + len(table_pages) + len(figure_pages),
        "table_label": thai_label_for_physical_page(table_physical),
        "figure_label": thai_label_for_physical_page(figure_physical),
        "symbols_label": thai_label_for_physical_page(symbols_physical),
        "main_entry_count": len(main_entries),
        "table_entry_count": len(table_entries),
        "figure_entry_count": len(figure_entries),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--combined-outline", required=True)
    parser.add_argument("--appendix-outline", required=True)
    parser.add_argument("--bibliography-pages", type=int, required=True)
    parser.add_argument("--output", required=True)
    build(parser.parse_args())
