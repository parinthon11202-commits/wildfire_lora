from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Wildfire_Lora\wildfire_lora-main")
TEMPLATE = Path(r"D:\เล่ม_project_จบ\เล่มปริญญานิพนธ์\Template ปริญญานิพนธ์\03 คำอธิบายสัญลักษณ์และคำย่อ.docx")
OUTPUT = ROOT / "บทคัดย่อพร้อมใส่เล่ม.docx"

FONT = "TH SarabunPSK"
SIZE = Pt(16)

TITLE = "ระบบแจ้งเตือนความเสี่ยงไฟป่าโดยใช้เครือข่ายสื่อสารระยะไกลในพื้นที่ไร้สัญญาณ"
AUTHORS = "ณัฐภัทร ไกรเมฆ และปรินทร ทองเลี่ยม"
YEAR = "พ.ศ. 2569"
DEPARTMENT = "วิศวกรรมคอมพิวเตอร์"
ADVISOR = "ผู้ช่วยศาสตราจารย์จักรพนธ์ อบมา"
CO_ADVISOR = "(ไม่มี)"

OBJECTIVE = (
    "ปริญญานิพนธ์นี้มีวัตถุประสงค์เพื่อพัฒนาต้นแบบตรวจวัดและแจ้งเตือนสัญญาณบ่งชี้ความเสี่ยงไฟป่า"
    "ในจุดที่ไม่มีอินเทอร์เน็ตโดยตรง โดยตรวจวัดอุณหภูมิ ความชื้นสัมพัทธ์ ควันหรืออนุภาค และตำแหน่ง "
    "แล้วส่งข้อมูลจาก Sensor Node จำนวน 2 ชุดผ่าน LoRa ความถี่ 433 เมกะเฮิรตซ์ไปยัง Gateway "
    "เพื่อจัดเก็บ ประเมินสถานะ และแสดงผลบนเว็บ Dashboard"
)

METHOD = (
    "Sensor Node รหัส NODE01 และ NODE02 ใช้บอร์ด LILYGO LoRa32 ร่วมกับเซนเซอร์ SHT31 "
    "Sharp GP2Y1014AU0F และ GPS NEO-6M เฟิร์มแวร์ตรวจสอบข้อมูล สร้างค่า Baseline "
    "และจำแนก 6 สถานะตั้งแต่ CALIBRATING ถึง SENSOR_FAULT จากควัน อุณหภูมิ และความชื้น "
    "Gateway ส่งข้อมูลผ่าน Wi-Fi และ HTTP ไปยัง Backend ซึ่งจัดเก็บข้อมูลใน MongoDB และให้บริการ"
    "เว็บ Dashboard การทดสอบในพื้นที่จำลองภายในเขตเมืองครอบคลุมเซนเซอร์ แพ็กเก็ตข้อมูล GPS "
    "ระยะสื่อสาร LoRa และระบบพลังงานจากแผงโซลาร์เซลล์"
)

RESULTS = (
    "Sensor Node ทั้งสองชุดเริ่มต้นอุปกรณ์ อ่าน SHT31 สร้างแพ็กเก็ต JSON และเข้าสู่สถานะ NORMAL "
    "หลังสร้าง Baseline 12 รอบในเวลาประมาณ 9 นาที NODE01 วัดอุณหภูมิ 24.26-37.66 °C และ"
    "ความชื้นสัมพัทธ์ 49.18-83.93% ส่วน NODE02 วัดได้ 25.68-26.01 °C และ 64.26-66.29% ตามลำดับ "
    "NODE01 ให้ค่าควันสูงสุด 4,095 raw "
    "ส่วน NODE02 ยังไม่มีผลทดสอบควัน ทั้งสองชุดตรวจพบการถอด SHT31 เป็น SENSOR_FAULT และรับพิกัด "
    "GPS ได้ภายใน 10 นาที การสื่อสาร LoRa สำเร็จ 98% ที่ระยะ 1.1-3.0 กิโลเมตร และ 95% ที่ระยะ 4.1 "
    "กิโลเมตร แต่ไม่พบสัญญาณที่ระยะ 5.0 กิโลเมตร วงจร MT3608 ให้แรงดัน 5.02 V เมื่อไม่มีโหลดและ "
    "4.96-5.04 V เมื่อมีโหลด ซึ่งผ่านเกณฑ์ 4.75-5.25 V ผลทดสอบตรรกะความเสี่ยง เส้นทางข้อมูล"
    "ครบระบบ และการทำงานต่อเนื่อง 24 ชั่วโมงยังไม่ครบ จึงยังสรุปความแม่นยำและความพร้อมใช้งาน"
    "ในพื้นที่จริงไม่ได้"
)

KEYWORDS = "ไฟป่า, LoRa, เครือข่ายเซนเซอร์ไร้สาย, ระบบเฝ้าระวัง, พื้นที่ไร้สัญญาณ"


def set_run_font(run, *, bold=False, size=SIZE):
    run.font.name = FONT
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")


def remove_all_body_content(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = False

    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), "thaiLetter")
    pg_num.set(qn("w:start"), "4")


def configure_header(section):
    for header in (section.header, section.first_page_header, section.even_page_header):
        hdr = header._element
        for child in list(hdr):
            hdr.remove(child)
        hdr.append(OxmlElement("w:p"))
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run_font(p.add_run("ง"))


def configure_normal_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = SIZE
    style.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)


def add_metadata(doc, label, value, *, value_size=SIZE):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.left_indent = Cm(4.0)
    fmt.first_line_indent = Cm(-4.0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.tab_stops.add_tab_stop(Cm(4.0))
    set_run_font(p.add_run(label))
    set_run_font(p.add_run("\t"))
    set_run_font(p.add_run(value), size=value_size)
    return p


def add_abstract_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(1.27)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(2)
    fmt.widow_control = True
    set_run_font(p.add_run(label + " " + text))
    return p


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build():
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    remove_all_body_content(doc)
    configure_section(doc.sections[0])
    configure_header(doc.sections[0])
    configure_normal_style(doc)

    add_metadata(doc, "หัวข้อปริญญานิพนธ์", TITLE, value_size=Pt(12))
    add_metadata(doc, "จัดทำโดย", AUTHORS)
    add_metadata(doc, "ปีที่ปริญญานิพนธ์สำเร็จ", YEAR)
    add_metadata(doc, "สาขาวิชา", DEPARTMENT)
    add_metadata(doc, "ประธานที่ปรึกษา", ADVISOR)
    add_metadata(doc, "ที่ปรึกษาร่วม", CO_ADVISOR)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run_font(heading.add_run("บทคัดย่อ"))

    add_abstract_paragraph(doc, "จุดประสงค์", OBJECTIVE)
    add_abstract_paragraph(doc, "วิธีการดำเนินงาน", METHOD)
    add_abstract_paragraph(doc, "ผลการดำเนินงาน", RESULTS)

    keywords = doc.add_paragraph()
    keywords.alignment = WD_ALIGN_PARAGRAPH.LEFT
    keywords.paragraph_format.space_before = Pt(3)
    keywords.paragraph_format.space_after = Pt(0)
    keywords.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run_font(keywords.add_run("คำสำคัญ\t:  "), bold=True)
    set_run_font(keywords.add_run(KEYWORDS))

    set_update_fields(doc)
    doc.core_properties.title = "บทคัดย่อระบบแจ้งเตือนความเสี่ยงไฟป่าโดยใช้เครือข่ายสื่อสารระยะไกลในพื้นที่ไร้สัญญาณ"
    doc.core_properties.subject = "บทคัดย่อปริญญานิพนธ์"
    doc.core_properties.keywords = KEYWORDS
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
