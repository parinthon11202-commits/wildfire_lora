import argparse
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document):
    body = document.element.body
    p_index = 0
    t_index = 0
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            yield ("P", p_index, paragraph)
            p_index += 1
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            yield ("T", t_index, table)
            t_index += 1


def run_flags(paragraph):
    flags = []
    for run in paragraph.runs:
        text = run.text.strip()
        if not text:
            continue
        color = run.font.color.rgb
        if color is not None and str(color).upper() not in {"000000", "00000A", "0A0000"}:
            flags.append(f"color={color}:{text}")
        if run.bold:
            flags.append(f"bold:{text}")
    return flags


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx")
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--end", type=int, default=10**9)
    parser.add_argument("--nonempty", action="store_true")
    args = parser.parse_args()

    document = Document(args.docx)
    print(f"FILE={Path(args.docx)}")
    print(f"PARAGRAPHS={len(document.paragraphs)} TABLES={len(document.tables)} "
          f"INLINE_SHAPES={len(document.inline_shapes)} SECTIONS={len(document.sections)}")

    block_index = 0
    for kind, local_index, obj in iter_blocks(document):
        if block_index < args.start:
            block_index += 1
            continue
        if block_index >= args.end:
            break
        if kind == "P":
            text = obj.text.replace("\n", " ").strip()
            if text or not args.nonempty:
                flags = run_flags(obj)
                print(f"B{block_index:04d} P{local_index:04d} [{obj.style.name}] {text}")
                if flags:
                    print("  RUNFLAGS " + " | ".join(flags))
        else:
            print(f"B{block_index:04d} T{local_index:03d} rows={len(obj.rows)} cols={len(obj.columns)}")
            for r_idx, row in enumerate(obj.rows):
                cells = []
                for cell in row.cells:
                    cell_text = " / ".join(p.text.replace("\n", " ").strip() for p in cell.paragraphs)
                    cells.append(cell_text)
                print(f"  R{r_idx:03d}: " + " || ".join(cells))
        block_index += 1

    print(f"BLOCKS_SCANNED_TO={block_index}")


if __name__ == "__main__":
    main()
