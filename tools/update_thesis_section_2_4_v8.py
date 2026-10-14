from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ZERO_WIDTH = re.compile(r"[\u200b\ufeff]")


def clean(text: str) -> str:
    return ZERO_WIDTH.sub("", text).strip()


def find_paragraph(doc: Document, needle: str, start: int = 0) -> int:
    for index in range(start, len(doc.paragraphs)):
        if needle in clean(doc.paragraphs[index].text):
            return index
    raise ValueError(f"paragraph not found: {needle}")


def replace_text_preserving_paragraph(paragraph, text: str) -> None:
    font_name = "TH SarabunPSK"
    font_size = Pt(16)
    if paragraph.runs:
        first_run = paragraph.runs[0]
        if first_run.font.name:
            font_name = first_run.font.name
        elif first_run._element.rPr is not None and first_run._element.rPr.rFonts is not None:
            font_name = first_run._element.rPr.rFonts.get(qn("w:eastAsia")) or font_name
        if first_run.font.size:
            font_size = first_run.font.size
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.font.name = font_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
    run.font.size = font_size


def insert_picture_before(caption, image_path: Path, width_inches: float = 5.77):
    picture_paragraph = caption.insert_paragraph_before()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(6)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    for doc_pr in picture_paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", "แผนภาพสถาปัตยกรรมซอฟต์แวร์และการไหลของข้อมูลในโครงงาน รุ่น 8")
    return picture_paragraph


def remove_existing_picture_between(doc: Document, start_index: int, caption_index: int) -> None:
    for paragraph in doc.paragraphs[start_index + 1 : caption_index]:
        if paragraph._p.xpath(".//a:blip"):
            paragraph._p.getparent().remove(paragraph._p)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("figure", type=Path)
    args = parser.parse_args()

    doc = Document(args.source)
    section_start = find_paragraph(doc, "เทคโนโลยีซอฟต์แวร์และการจัดการข้อมูล")
    intro_index = section_start + 1
    caption_index = find_paragraph(doc, "รูปที่ 2.13", intro_index + 1)

    intro = doc.paragraphs[intro_index]
    replace_text_preserving_paragraph(
        intro,
        "ระบบซอฟต์แวร์ของโครงงานประกอบด้วยเฟิร์มแวร์ Sensor Node สำหรับอ่านค่า SHT31 และ Sharp "
        "ตรวจสุขภาพเซนเซอร์ ประเมินสถานะตามกฎรุ่น 8 และสร้างแพ็กเก็ต JSON ส่งผ่าน LoRa; "
        "เฟิร์มแวร์ Gateway สำหรับรับและตรวจแพ็กเก็ต ส่ง ACK หรือคำสั่งกลับไปยังโหนด "
        "และส่งต่อข้อมูลผ่าน Wi-Fi/HTTPS; Backend สำหรับตรวจ API key รูปแบบแพ็กเก็ต รุ่นกฎ "
        "ข้อมูลซ้ำและข้อมูลเก่า แล้วจัดเก็บค่าตรวจวัดกับสถานะที่โหนดตัดสิน รวมทั้งจัดการเหตุการณ์แจ้งเตือนและคิวคำสั่ง; "
        "MongoDB สำหรับจัดเก็บข้อมูล; และ Dashboard สำหรับแสดงสถานะ ค่าตรวจวัด แผนที่ กราฟ และเหตุการณ์ผ่าน HTTP API "
        "โดย Gateway และ Backend ไม่คำนวณความเสี่ยงซ้ำ ดังแสดงในรูปที่ 2.13",
    )
    intro.style = doc.styles["Normal"]
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.first_line_indent = Inches(0.5)
    intro.paragraph_format.space_after = Pt(6)

    # Re-fetch after text changes, then ensure exactly one figure exists before the caption.
    caption_index = find_paragraph(doc, "รูปที่ 2.13", intro_index + 1)
    remove_existing_picture_between(doc, intro_index, caption_index)
    caption = doc.paragraphs[find_paragraph(doc, "รูปที่ 2.13", intro_index + 1)]
    insert_picture_before(caption, args.figure)

    replacements = {
        "เฟิร์มแวร์ Arduino และรูปแบบ JSON": "เฟิร์มแวร์ Arduino และรูปแบบ JSON",
        "Arduino Framework สำหรับ ESP32": (
            "Arduino Framework สำหรับ ESP32 ทำให้พัฒนาเฟิร์มแวร์ด้วย C/C++ และใช้ไลบรารีสำหรับ SHT31, LoRa, GPS, Wi-Fi "
            "และการจัดเก็บค่าคงอยู่ใน Preferences ได้ [8] ในโครงงาน Sensor Node อ่าน SHT31 และแรงดันจาก Sharp "
            "ตรวจความสมเหตุสมผลของข้อมูล เลือกค่ามัธยฐานของ Sharp สามตัวอย่าง ประเมิน NORMAL, WATCH หรือ WARNING "
            "ตามกฎรุ่น 8 และกำหนด SENSOR_FAULT เมื่อเซนเซอร์ผิดปกติ ก่อนสร้างแพ็กเก็ตและจัดการรอบประหยัดพลังงาน"
        ),
        "JavaScript ObjectNotation": (
            "JavaScript Object Notation (JSON) เป็นรูปแบบข้อมูลข้อความตามมาตรฐาน RFC 8259 [9] ระบบใช้ JSON แบบย่อในช่องทาง LoRa "
            "โดยแพ็กเก็ต Sensor รุ่น 8 มีฟิลด์หลัก t, id, q, sid, ri, st, rv, at, h, pm และ sh "
            "และไม่มีคะแนนความเสี่ยงหรือรหัสเหตุผลแบบบิต Gateway เติมข้อมูลคุณภาพสัญญาณ RSSI/SNR ก่อนส่งต่อผ่าน HTTP ไปยัง Backend"
        ),
        "การตรวจคำขอ และการตอบกลับ": (
            "การตรวจคำขอ และการตอบกลับ [11], [12] ในโครงงาน Backend ทำหน้าที่ตรวจ API key โครงสร้างแพ็กเก็ตและ rv = 8 "
            "สร้างเอกลักษณ์ข้อมูลเพื่อป้องกันข้อมูลซ้ำ ตรวจข้อมูลเก่า บันทึกค่าตรวจวัดกับสถานะที่ Sensor Node ส่งมา "
            "จัดการเหตุการณ์แจ้งเตือน และจัดการคิวคำสั่ง GPS โดยไม่คำนวณเกณฑ์ความเสี่ยงซ้ำ"
        ),
        "MongoDB จัดเก็บข้อมูลเป็น BSON": (
            "MongoDB จัดเก็บข้อมูลเป็น BSON Document ภายใน Collection ซึ่งมีโครงสร้างใกล้เคียงวัตถุ JSON และรองรับข้อมูลซ้อนกับอาร์เรย์ [13] "
            "รูปแบบดังกล่าวเหมาะกับข้อมูลเซนเซอร์ที่ประกอบด้วยค่าตรวจวัด สถานะที่โหนดตัดสิน รุ่นกฎ สุขภาพเซนเซอร์ พิกัด "
            "คุณภาพสัญญาณ และแพ็กเก็ตต้นทางในเอกสารเดียว"
        ),
        "React ใช้สร้างส่วนติดต่อผู้ใช้": (
            "React ใช้สร้างส่วนติดต่อผู้ใช้แบบองค์ประกอบ [14] ส่วน TypeScript เพิ่มการตรวจชนิดข้อมูล ทำให้โครงสร้างค่าที่รับจาก API ชัดเจนขึ้น [27] "
            "Leaflet ใช้แสดงแผนที่และหมุดตำแหน่งโหนด [15] ขณะที่ Recharts ใช้สร้างกราฟแนวโน้มอุณหภูมิ ความชื้น "
            "และค่าความเข้มข้นอนุภาคโดยประมาณจากข้อมูลตามเวลา [16]"
        ),
        "CloudflareTunnel สร้างการเชื่อมต่อ": (
            "Cloudflare Tunnel สร้างการเชื่อมต่อขาออกจากเซิร์ฟเวอร์ภายในไปยังเครือข่าย Cloudflare จึงไม่จำเป็นต้องเปิดพอร์ตขาเข้าจากอินเทอร์เน็ตโดยตรง "
            "การตั้งค่าของโครงงานแยก Dashboard แบบอ่านอย่างเดียวที่ wildfire.nattaphat.me ออกจากส่วนผู้ดูแลที่ admin.nattaphat.me "
            "ซึ่งป้องกันด้วย Cloudflare Access [17] อย่างไรก็ตาม Backend ยังคงตรวจสิทธิ์ของคำขอที่เปลี่ยนข้อมูลและไม่พึ่งการป้องกันที่ชั้นหน้าเพียงอย่างเดียว"
        ),
        "Telegram Bot API เป็นส่วนติดต่อ": (
            "Telegram Bot API เป็นส่วนติดต่อแบบ HTTP ที่ช่วยให้ระบบภายนอกส่งข้อความผ่านบอตไปยังผู้ใช้งานหรือกลุ่มสนทนา [28] "
            "ในโครงงาน Backend ใช้ส่งข้อความเมื่อเกิดหรือเปลี่ยนเหตุการณ์ WATCH, WARNING หรือ SENSOR_FAULT "
            "โดยข้อความประกอบด้วยรหัสโหนด สถานะ ค่าตรวจวัด ตำแหน่ง และเวลา ทั้งนี้ Bot Token และข้อมูลผู้รับจัดเก็บไว้ฝั่งเซิร์ฟเวอร์ "
            "และไม่เปิดเผยในเฟิร์มแวร์หรือส่วนติดต่อผู้ใช้"
        ),
    }

    section_end = find_paragraph(doc, "แนวทางที่ใช้ในการประเมินสถานะของโครงงาน", section_start)
    for paragraph in doc.paragraphs[section_start:section_end]:
        normalized = clean(paragraph.text)
        for prefix, new_text in replacements.items():
            if normalized.startswith(prefix):
                replace_text_preserving_paragraph(paragraph, new_text)
                break

    set_update_fields(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
