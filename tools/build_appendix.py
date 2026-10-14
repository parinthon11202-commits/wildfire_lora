from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Wildfire_Lora\wildfire_lora-main")
TEMPLATE = Path(r"D:\เล่ม_project_จบ\เล่มปริญญานิพนธ์\Template ปริญญานิพนธ์\10 ภาคผนวก.docx")
SOURCE_MD = ROOT / ".tmp" / "appendix_artifact.md"
OUTPUT = ROOT / "ภาคผนวก ก คู่มือการติดตั้งและใช้งานระบบ.docx"

THAI_FONT = "TH Sarabun New"
BODY_SIZE = Pt(16)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=BODY_SIZE, bold=None, italic=None, color=None, font=THAI_FONT):
    run.font.name = font
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_defaults(paragraph, *, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                           first_line=True, after=Pt(0), keep_next=False):
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = after
    fmt.widow_control = True
    fmt.keep_with_next = keep_next
    fmt.first_line_indent = Inches(0.5) if first_line else None


INLINE_RE = re.compile(r"(`[^`]+`)")


def add_inline_runs(paragraph, text: str, *, bold=False, size=BODY_SIZE):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=Pt(12.5), font="Consolas")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold)


def add_heading(doc, text: str):
    p = doc.add_paragraph(style="Appendix Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(10)
    fmt.space_after = Pt(2)
    fmt.line_spacing = 1.0
    fmt.keep_with_next = True
    fmt.keep_together = True
    run = p.add_run(text)
    set_run_font(run, size=Pt(18), bold=True)
    return p


def add_body(doc, text: str):
    p = doc.add_paragraph()
    set_paragraph_defaults(p)
    add_inline_runs(p, text)
    return p


def add_list_item(doc, text: str, marker: str):
    p = doc.add_paragraph()
    set_paragraph_defaults(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    run = p.add_run(f"{marker} ")
    set_run_font(run)
    add_inline_runs(p, text)
    return p


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return None
    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_pr = table._tbl.tblPr
    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "808080")
    tbl_layout = table_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        table_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "autofit")

    for r_idx, row_values in enumerate(rows):
        table_row = table.rows[r_idx]
        for c_idx in range(columns):
            value = row_values[c_idx] if c_idx < len(row_values) else ""
            cell = table_row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "D9E2F3")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            fmt = p.paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.0
            fmt.widow_control = True
            add_inline_runs(p, value, bold=(r_idx == 0), size=Pt(14))
        tr_pr = table_row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if r_idx == 0:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.5
    set_run_font(spacer.add_run(""), size=Pt(6))
    return table


def parse_markdown_body(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    body_started = False
    items = []
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        if raw.startswith("### "):
            body_started = True
            items.append(("heading", raw[4:].strip()))
            i += 1
            continue
        if not body_started:
            i += 1
            continue
        if not raw.strip():
            i += 1
            continue
        if raw.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = []
            for idx, line in enumerate(table_lines):
                cells = [c.strip() for c in line.strip("|").split("|")]
                if idx == 1 and all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    continue
                parsed.append(cells)
            items.append(("table", parsed))
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", raw)
        if numbered:
            items.append(("numbered", (numbered.group(1) + ".", numbered.group(2))))
            i += 1
            continue
        bullet = re.match(r"^-\s+(.*)$", raw)
        if bullet:
            items.append(("bullet", bullet.group(1)))
            i += 1
            continue
        items.append(("body", raw.strip()))
        i += 1
    return items


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = THAI_FONT
    normal.font.size = BODY_SIZE
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), THAI_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), THAI_FONT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    try:
        heading = doc.styles["Appendix Heading 1"]
    except KeyError:
        heading = doc.styles.add_style("Appendix Heading 1", WD_STYLE_TYPE.PARAGRAPH)
        heading.base_style = normal
    heading.font.name = THAI_FONT
    heading.font.size = Pt(18)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), THAI_FONT)
    heading._element.rPr.rFonts.set(qn("w:cs"), THAI_FONT)
    p_pr = heading._element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), "0")


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_section_vertical_alignment(section, value: str) -> None:
    sect_pr = section._sectPr
    valign = sect_pr.find(qn("w:vAlign"))
    if valign is None:
        valign = OxmlElement("w:vAlign")
        sect_pr.append(valign)
    valign.set(qn("w:val"), value)


def build():
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    configure_styles(doc)

    # Keep the original section-break paragraph and header/footer definitions,
    # but remove all sample AR-project body content and images.
    section_break_idx = None
    for idx, p in enumerate(doc.paragraphs):
        p_pr = p._p.pPr
        if p_pr is not None and p_pr.sectPr is not None:
            section_break_idx = idx
            break
    if section_break_idx is None:
        raise RuntimeError("The appendix template has no section break to preserve.")

    original_paragraphs = list(doc.paragraphs)
    break_paragraph = original_paragraphs[section_break_idx]
    for p in original_paragraphs[section_break_idx + 1:]:
        remove_paragraph(p)
    for p in original_paragraphs[:section_break_idx]:
        if p is not original_paragraphs[0]:
            remove_paragraph(p)

    cover = original_paragraphs[0]
    clear_paragraph(cover)
    cover.style = doc.styles["Normal"]
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.first_line_indent = None
    cover.paragraph_format.line_spacing = 1.15
    cover.paragraph_format.space_before = Pt(0)
    cover.paragraph_format.space_after = Pt(0)
    r = cover.add_run("ภาคผนวก ก")
    set_run_font(r, size=Pt(20), bold=True)
    r.add_break()
    r.add_break()
    r2 = cover.add_run("คู่มือการติดตั้งและใช้งานระบบตรวจสอบและแจ้งเตือนความเสี่ยงไฟป่า")
    set_run_font(r2, size=Pt(18), bold=True)
    r2.add_break()
    r3 = cover.add_run("ผ่านการสื่อสาร LoRa")
    set_run_font(r3, size=Pt(18), bold=True)

    clear_paragraph(break_paragraph)
    break_paragraph.paragraph_format.space_after = Pt(0)

    sections = doc.sections
    if len(sections) < 2:
        raise RuntimeError("The appendix template must retain two sections.")
    set_section_vertical_alignment(sections[0], "center")
    sections[0].different_first_page_header_footer = True
    set_section_vertical_alignment(sections[1], "top")
    for section in sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.5)
        section.bottom_margin = Inches(1.0)
        section.header_distance = Inches(1.0)
        section.footer_distance = Inches(0.492)

    items = parse_markdown_body(SOURCE_MD)
    for kind, value in items:
        if kind == "heading":
            add_heading(doc, value)
        elif kind == "table":
            add_table(doc, value)
        elif kind == "numbered":
            marker, text = value
            add_list_item(doc, text, marker)
        elif kind == "bullet":
            add_list_item(doc, value, "•")
        else:
            add_body(doc, value)

    # Use paragraph spacing, not empty spacer paragraphs, at the document end.
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        p = doc.paragraphs[-1]
        p_pr = p._p.pPr
        if p_pr is not None and p_pr.sectPr is not None:
            break
        remove_paragraph(p)

    set_update_fields(doc)
    doc.core_properties.title = "ภาคผนวก ก คู่มือการติดตั้งและใช้งานระบบตรวจสอบและแจ้งเตือนความเสี่ยงไฟป่า"
    doc.core_properties.subject = "คู่มือการติดตั้ง ตั้งค่า ใช้งาน และแก้ไขปัญหาเบื้องต้น"
    doc.core_properties.keywords = "ไฟป่า, LoRa, Sensor Node, Gateway, Backend, Dashboard"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
