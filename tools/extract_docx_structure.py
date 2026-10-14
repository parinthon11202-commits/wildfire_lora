import argparse
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def run_props(run):
    rpr = run._element.rPr
    fonts = {}
    if rpr is not None and rpr.rFonts is not None:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            value = rpr.rFonts.get(qn(f"w:{key}"))
            if value:
                fonts[key] = value
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": bool(run.underline) if run.underline is not None else None,
        "size_pt": run.font.size.pt if run.font.size else None,
        "font": run.font.name,
        "fonts": fonts,
    }


def paragraph_info(paragraph, index=None):
    pf = paragraph.paragraph_format
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "left_indent_pt": pf.left_indent.pt if pf.left_indent else None,
        "first_line_indent_pt": pf.first_line_indent.pt if pf.first_line_indent else None,
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": float(pf.line_spacing) if isinstance(pf.line_spacing, (int, float)) else None,
        "runs": [run_props(run) for run in paragraph.runs if run.text],
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()

    doc = Document(args.input)
    result = {
        "path": str(Path(args.input).resolve()),
        "paragraphs": [paragraph_info(p, i) for i, p in enumerate(doc.paragraphs)],
        "tables": [],
        "sections": [],
        "inline_shapes": len(doc.inline_shapes),
        "styles": [],
    }

    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                cells.append({
                    "row": ri,
                    "col": ci,
                    "text": "\n".join(p.text for p in cell.paragraphs),
                    "paragraphs": [paragraph_info(p) for p in cell.paragraphs],
                })
            rows.append(cells)
        result["tables"].append({
            "index": ti,
            "style": table.style.name if table.style else None,
            "rows": rows,
        })

    for si, section in enumerate(doc.sections):
        result["sections"].append({
            "index": si,
            "width_mm": section.page_width.mm,
            "height_mm": section.page_height.mm,
            "top_margin_mm": section.top_margin.mm,
            "bottom_margin_mm": section.bottom_margin.mm,
            "left_margin_mm": section.left_margin.mm,
            "right_margin_mm": section.right_margin.mm,
            "header_distance_mm": section.header_distance.mm,
            "footer_distance_mm": section.footer_distance.mm,
            "header": [paragraph_info(p) for p in section.header.paragraphs],
            "footer": [paragraph_info(p) for p in section.footer.paragraphs],
        })

    for style in doc.styles:
        if style.type == 1:
            font = style.font
            pf = style.paragraph_format
            result["styles"].append({
                "name": style.name,
                "font": font.name,
                "size_pt": font.size.pt if font.size else None,
                "bold": font.bold,
                "italic": font.italic,
                "alignment": int(pf.alignment) if pf.alignment is not None else None,
                "left_indent_pt": pf.left_indent.pt if pf.left_indent else None,
                "right_indent_pt": pf.right_indent.pt if pf.right_indent else None,
                "first_line_indent_pt": pf.first_line_indent.pt if pf.first_line_indent else None,
                "space_before_pt": pf.space_before.pt if pf.space_before else None,
                "space_after_pt": pf.space_after.pt if pf.space_after else None,
                "line_spacing": float(pf.line_spacing) if isinstance(pf.line_spacing, (int, float)) else None,
                "keep_with_next": pf.keep_with_next,
                "keep_together": pf.keep_together,
                "page_break_before": pf.page_break_before,
            })

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
